import torch
import torch.nn as nn
import os.path as osp
import wandb
from types import SimpleNamespace
from torch.optim import Adam, lr_scheduler
from torch import Tensor
from torch.nn.functional import softmax, relu
import pandas as pd
import re
import ast
import openpyxl
from omegaconf import OmegaConf
from torchmetrics import ConfusionMatrix

from utils import get_device, BinClf_Evaluator, MultiClf_Evaluator, get_loss_function, get_init_key_frequency, mkdir_if_missing, get_current_ensemble_classes, rename_keys, check_feature_distribution, plot_key_matching_heatmap, get_current_eval_dataloader
from models.conch import create_model_from_pretrained
from models import freeze_weight, set_tunable_v, EarlyStopping, AlignMIL, QPMIL_VL
from dataset import get_data_loaders, get_sids
from utils.tools import get_null_space, plot_confusion_matrix


class Manager(object):
    def __init__(self, cfg):
        self.cfg = cfg
        # AlignMIL (ours) or the QPMIL_VL baseline; PCFA and the null space update are only used by AlignMIL
        self.model_name = cfg.get('model', 'AlignMIL')
        assert self.model_name in ['AlignMIL', 'QPMIL_VL'], 'Invalid model: {}'.format(self.model_name)

        self.device = get_device(cfg['cuda_id'])

        if cfg['base_model_arch'] == 'CONCH':
            self.base_model, _ = create_model_from_pretrained("conch_ViT-B-16", checkpoint_path=cfg['conch_ckpt_path'], device=self.device)
            self.base_model.eval()
            self.base_model.dtype = self.base_model.logit_scale.dtype
            self.embedding_dim = self.base_model.text.ln_final.weight.shape[0] # D_e: embedding dimension of E_txt
            self.feature_dim = self.base_model.visual.proj_contrast.shape[1] # D_f: output feature dimension of encoder E_img and E_txt
        else:
            raise NotImplementedError("Please specify a valid architecture.")
        freeze_weight(self.base_model, cfg['base_model_arch'])
        self.dtype = self.base_model.dtype

        self.writer = None

        self.model = None
        self.optimizer = None
        self.lr_scheduler = None
        self.early_stop = None

        self.b_evaluator = BinClf_Evaluator()
        self.m_evaluator = MultiClf_Evaluator()

        self.loss_function = get_loss_function(cfg['loss_function'])
        print('[setup] loss function:', cfg['loss_function'])

        self.current_dataset = None
        self.current_save_result_dir = None
        self.current_ensemble_classes = None

        self.data_loader = get_data_loaders(cfg)
        self.sids = get_sids(self.data_loader)

        # Parameters to be trained: Prototype Pool (key, prompt)
        if cfg['opt_name'] == 'adam':
            self.key = nn.ParameterList([nn.Parameter(0.02 * torch.randn(1, self.feature_dim, dtype=self.dtype, device=self.device)) for _ in range(cfg['pool_size'])])
            if self.model_name == 'AlignMIL': # learnable prototype vectors
                self.prompt = nn.ParameterList([nn.Parameter(0.02 * torch.randn(1, self.feature_dim, dtype=self.dtype, device=self.device)) for _ in range(cfg['pool_size'])])
            else: # learnable prompts encoded by the language encoder
                self.prompt = nn.ParameterList([nn.Parameter(0.02 * torch.randn(1, cfg['prompt_length'], self.embedding_dim, dtype=self.dtype, device=self.device)) for _ in range(cfg['pool_size'])])
            print('[setup] Parameters to be trained:\nPrototype Pool\nkey:\n{}\nprompt:\n{}'.format(self.key, self.prompt))
        else:
            raise NotImplementedError('Invalid optimizer')

        # Parameters to be trained: Tunable Vector
        self.tunable_v = nn.ParameterList([nn.Parameter(torch.zeros(class_num, self.feature_dim, dtype=self.dtype, device=self.device)) for class_num in cfg['dataset_subtype_num']])
        print('Tunable Vector\n{}'.format(self.tunable_v))

        self.train_key_frequency = get_init_key_frequency(cfg) # {'dataset_name': [[0, 0, ...], []]} len(v)=epochs, v[0]=tensor([pool_size])
        self.val_key_frequency = get_init_key_frequency(cfg)
        self.dataset2key_idx = {} # {'dataset_name': [0, 1, ...]} the matched key/prompt idx of the dataset

        self.sum_test_acc = 0
        self.test_acc = []
        self.cfg["seen_task_dict"] = {} # {'dataset_name': idx} the idx of the task
        self.cfg["idx2dataset"] = {} # {idx: 'dataset_name'} the dataset name of the task

        # null space of each class: phi (basis), alpha (update coefficients), fb (query features of its task)
        self.phi = {}
        self.alpha = {}
        self.fb = {}

    def incre_train(self):
        for dataset_name in self.cfg['dataset_names']:
            self.cfg['task_num'] += 1
            print('[train] Task: {}, Dataset: {}'.format(self.cfg['task_num'], dataset_name))
            if dataset_name in self.cfg["seen_task_dict"]:
                duplicate_flag = True
            else:
                self.cfg["seen_task_dict"][dataset_name] = len(self.cfg["seen_task_dict"])
                self.cfg["idx2dataset"][len(self.cfg["seen_task_dict"])-1] = dataset_name
                duplicate_flag = False

            self.current_dataset = dataset_name
            self.current_dataset_idx = self.cfg["seen_task_dict"][dataset_name]
            self.current_save_result_dir = osp.join(self.cfg['save_result_dir'], dataset_name+'-task'+str(self.cfg['task_num']))
            mkdir_if_missing(self.current_save_result_dir)
            self.current_ensemble_classes = get_current_ensemble_classes(self.cfg, self.cfg["seen_task_dict"].keys())
            print('[infor] Current dataset: {}'.format(self.current_dataset))
            print('[infor] Current save result dir: {}'.format(self.current_save_result_dir))

            # wandb writer
            # save_result_dir: e.g. "results/ex_cl_fw/2025-03-04/train-data_split_seed_1"
            project_name = "AlignMIL_" + self.cfg['save_result_dir'].split('/')[1]
            group_name = self.cfg['save_result_dir'].split('/')[2]
            run_name = self.cfg['save_result_dir'].split('/')[-1] + '-' + dataset_name + '-task' + str(self.cfg['task_num'])
            wandb_mode = self.cfg.get('wandb', 'online')
            if wandb_mode is True: wandb_mode = 'online'
            if wandb_mode is False: wandb_mode = 'disabled'
            wandb_config = OmegaConf.to_container(self.cfg, resolve=True, throw_on_missing=True)
            self.writer = wandb.init(project=project_name, group=group_name, name=run_name, dir=self.current_save_result_dir, config=wandb_config, reinit=True, mode=wandb_mode)
            print('[setup] Wandb writer construction completed')

            # model
            if self.model_name == 'AlignMIL':
                set_tunable_v(self.tunable_v, self.current_dataset_idx) # Only the tunable vectors of the current dataset classes are trained.
                self.model = AlignMIL(self.cfg, self.base_model, self.device, self.key, self.prompt, self.tunable_v, self.current_ensemble_classes, self.train_key_frequency)
            else:
                self.model = QPMIL_VL(self.cfg, self.base_model, self.device, self.key, self.prompt, self.tunable_v, self.current_ensemble_classes, self.train_key_frequency)
            print('[setup] {} construction completed'.format(self.model_name))

            # optimizer
            if self.cfg['opt_name'] == 'adam':
                cfg_optimizer = SimpleNamespace(opt=self.cfg['opt_name'], weight_decay=self.cfg['adam_weight_decay'],
                                                lr=self.cfg['adam_lr'], opt_eps=self.cfg['adam_eps'], opt_betas=None,
                                                momentum=None)
                self.optimizer = Adam(self.model.parameters(), lr=cfg_optimizer.lr,
                                      weight_decay=cfg_optimizer.weight_decay, eps=cfg_optimizer.opt_eps)
            else:
                raise NotImplementedError('Invalid optimizer')
            print("[setup] optimizer:", cfg_optimizer)

            # lr_scheduler
            cfg_lr_scheduler = SimpleNamespace(name=self.cfg['lrs_name'], mode=self.cfg['lrs_mode'],
                                               factor=self.cfg['lrs_factor'],
                                               patience=self.cfg['lrs_patience'][self.cfg['task_num'] - 1],
                                               threshold=self.cfg['lrs_threshold'],
                                               threshold_mode=self.cfg['lrs_threshold_mode'])
            if cfg_lr_scheduler.name == 'ReduceLROnPlateau':
                self.lr_scheduler = lr_scheduler.ReduceLROnPlateau(self.optimizer, mode=cfg_lr_scheduler.mode,
                                                                   factor=cfg_lr_scheduler.factor,
                                                                   patience=cfg_lr_scheduler.patience,
                                                                   threshold=cfg_lr_scheduler.threshold,
                                                                   threshold_mode=cfg_lr_scheduler.threshold_mode)
            else:
                raise NotImplementedError('Invalid lr scheduler')
            print("[setup] lr_scheduler:", cfg_lr_scheduler)

            # early_stop
            self.early_stop = EarlyStopping(warmup=self.cfg['es_warmup'],
                                            patience=self.cfg['es_patience'][self.cfg['task_num'] - 1],
                                            verbose=self.cfg['es_verbose'], threshold=self.cfg['es_threshold'])
            print(f"[setup] early_stop: warmup={self.cfg['es_warmup']}, patience={self.cfg['es_patience'][self.cfg['task_num'] - 1]}, verbose={self.cfg['es_verbose']}, threshold={self.cfg['es_threshold']}")

            # run training
            self._run_training()

            if self.model_name == 'AlignMIL':
                self.cal_null_space()

            # evaluate at val/test dataset
            eval_result_dir = osp.join(self.current_save_result_dir, 'eval_results')
            mkdir_if_missing(eval_result_dir)
            current_eval_dataloader = get_current_eval_dataloader(self.data_loader, self.cfg["seen_task_dict"].keys())
            self._eval_all(current_eval_dataloader, eval_result_dir)
            print('[eval at val/test dataset] Path where the evaluation result is saved:', eval_result_dir)

            # calculate masked metrics at val/test dataset
            masked_eval_result_dir = eval_result_dir
            self._eval_masked_metrics(current_eval_dataloader, masked_eval_result_dir)
            print('[cal masked metrics at val/test dataset] Path where the evaluation result is saved:', masked_eval_result_dir)

            # save the enhanced class feature and prototype features
            enhanced_class_feature = self.model.get_enhanced_class_feature()
            torch.save(enhanced_class_feature.detach().cpu(), osp.join(self.current_save_result_dir, 'enhanced_class_feature.pth'))
            prototype_features_list = []
            for start_idx in range(0, self.cfg['pool_size'], self.cfg['match_size']):
                end_idx = min(start_idx + self.cfg['match_size'], self.cfg['pool_size'])
                batch_indices = torch.tensor(range(start_idx, end_idx))
                prototype_features = self.model.get_prototype_features(batch_indices)
                prototype_features_list.append(prototype_features)
            torch.save(torch.cat(prototype_features_list, dim=0).detach().cpu(), osp.join(self.current_save_result_dir, 'prototype_features.pth'))

        average_test_acc = self.sum_test_acc / len(self.cfg['dataset_names'])
        print('[result] Average test acc: {:.6f}'.format(average_test_acc))
        wandb.log({'result/avg_test_acc': average_test_acc})

        print("dataset2key_idx", self.dataset2key_idx.values())

        if self.cfg['data_split_seed'] == self.cfg['total_fold']:
            metrics_dir = osp.join(osp.dirname(osp.dirname(self.current_save_result_dir)), 'metrics')
            mkdir_if_missing(metrics_dir)

            for fold_num in range(1, self.cfg['total_fold'] + 1):
                with open(osp.join(metrics_dir, 'test_acc.txt'), 'a') as f:
                    read_path = osp.join(re.sub(r'_\d+$', '_' + str(fold_num), osp.dirname(self.current_save_result_dir)), 'metrics/test_acc.txt')
                    with open(read_path, 'r') as file_read:
                        f.write(file_read.read() + '\n')

                with open(osp.join(metrics_dir, 'val_acc.txt'), 'a') as f:
                    read_path = osp.join(re.sub(r'_\d+$', '_' + str(fold_num), osp.dirname(self.current_save_result_dir)), 'metrics/val_acc.txt')
                    with open(read_path, 'r') as file_read:
                        f.write(file_read.read() + '\n')

                with open(osp.join(metrics_dir, 'test_mask_acc.txt'), 'a') as f:
                    read_path = osp.join(re.sub(r'_\d+$', '_' + str(fold_num), osp.dirname(self.current_save_result_dir)), 'metrics/test_mask_acc.txt')
                    with open(read_path, 'r') as file_read:
                        f.write(file_read.read() + '\n')

                with open(osp.join(metrics_dir, 'test_mask_auc.txt'), 'a') as f:
                    read_path = osp.join(re.sub(r'_\d+$', '_' + str(fold_num), osp.dirname(self.current_save_result_dir)), 'metrics/test_mask_auc.txt')
                    with open(read_path, 'r') as file_read:
                        f.write(file_read.read() + '\n')

            result_metric_list = []
            for fold_num in range(1, self.cfg['total_fold'] + 1):
                read_path = osp.join(re.sub(r'_\d+$', '_' + str(fold_num), osp.dirname(self.current_save_result_dir)), 'metrics/test_acc.txt')
                with open(read_path, 'r') as file_read:
                    result_metric_list.extend(ast.literal_eval(file_read.read()))
            for fold_num in range(1, self.cfg['total_fold'] + 1):
                read_path = osp.join(re.sub(r'_\d+$', '_' + str(fold_num), osp.dirname(self.current_save_result_dir)), 'metrics/test_mask_acc.txt')
                with open(read_path, 'r') as file_read:
                    result_metric_list.append(ast.literal_eval(file_read.read()))
            for fold_num in range(1, self.cfg['total_fold'] + 1):
                read_path = osp.join(re.sub(r'_\d+$', '_' + str(fold_num), osp.dirname(self.current_save_result_dir)), 'metrics/test_mask_auc.txt')
                with open(read_path, 'r') as file_read:
                    result_metric_list.append(ast.literal_eval(file_read.read()))
            for fold_num in range(1, self.cfg['total_fold'] + 1):
                read_path = osp.join(re.sub(r'_\d+$', '_' + str(fold_num), osp.dirname(self.current_save_result_dir)), 'metrics/val_acc.txt')
                with open(read_path, 'r') as file_read:
                    result_metric_list.append(ast.literal_eval(file_read.read()))

            wb = openpyxl.load_workbook(self.cfg['eval_template_path'])
            ws = wb.active
            ws.cell(row=2, column=2, value=osp.dirname(osp.dirname(self.current_save_result_dir)).split('/')[-1])
            for i in range(len(self.cfg['dataset_names'])):
                ws.cell(row=3, column=i+5, value=self.cfg['dataset_names'][i])
            for row_index, row in enumerate(result_metric_list, start=4):
                for col_index, value in enumerate(row, start=5):
                    ws.cell(row=row_index, column=col_index, value=value)
            wb.save(osp.join(metrics_dir, osp.dirname(osp.dirname(self.current_save_result_dir)).split('/')[-1] + ' result.xlsx'))
            wb.close()

    def _eval_masked_metrics(self, data_loaders, masked_eval_result_dir):
        if self.current_dataset == self.cfg['dataset_names'][-1]:
            test_mask_acc = []
            test_mask_auc = []

        for k, data_loader in data_loaders.items():
            cltor = self.eval_model(self.model, data_loader)

            for k_cltor, v_cltor in cltor.items():
                label_shift = self.cfg['dataset_label_shift'][self.cfg['dataset_names'].index(k.split('/')[0])]
                v_cltor['y_hat'] = v_cltor['y_hat'][:, label_shift:label_shift + self.cfg['dataset_subtype_num'][self.cfg['dataset_names'].index(k.split('/')[0])]] # mask irrelevant logits
                if_binary = v_cltor['y_hat'].shape[1] == 2
                eval_results = self._eval_and_print(v_cltor, name=k + '/masked/' + k_cltor, if_binary=if_binary)

                if self.current_dataset == self.cfg['dataset_names'][-1]:
                    if k.split('/')[-1] == 'test':
                        mask_acc = eval_results['acc@mid'] if if_binary else eval_results['acc']
                        mask_acc = round(float(mask_acc), 6)
                        test_mask_acc.append(mask_acc)

                        if if_binary: # note: can't calculate auc when doing multi-classification
                            mask_auc = eval_results['auc']
                            mask_auc = round(float(mask_auc), 6)
                            test_mask_auc.append(mask_auc)

                sids = self._get_unique_sids(k, v_cltor['idx'])
                path_save_pred = masked_eval_result_dir + '/masked-' + k.replace('/', '-') + '.csv'
                self._save_prediction_clf(sids, v_cltor['y'], v_cltor['y_hat'], path_save_pred, binary=if_binary)

        if self.current_dataset == self.cfg['dataset_names'][-1]:
            metrics_dir = osp.join(osp.dirname(self.current_save_result_dir), 'metrics')
            mkdir_if_missing(metrics_dir)

            with open(osp.join(metrics_dir, 'test_mask_acc.txt'), 'w') as f:
                f.write(str(test_mask_acc))

            with open(osp.join(metrics_dir, 'test_mask_auc.txt'), 'w') as f:
                f.write(str(test_mask_auc))

    def _eval_all(self, data_loaders, eval_result_dir):
        # data_loaders: dict like {'camelyon16/val': loader1, 'camelyon16/test': loader2, ...}
        current_test_acc = []

        if self.current_dataset == self.cfg['dataset_names'][-1]:
            val_acc = []

        # confusion matrix init
        n_seen_classes = len(self.current_ensemble_classes['count'])
        print(f"[Confusion Matrix] Initializing for {n_seen_classes} classes.")

        if n_seen_classes == 2:
            conf_mat_val = ConfusionMatrix(task='binary', num_classes=n_seen_classes)
            conf_mat_test = ConfusionMatrix(task='binary', num_classes=n_seen_classes)
        else:
            conf_mat_val = ConfusionMatrix(task='multiclass', num_classes=n_seen_classes)
            conf_mat_test = ConfusionMatrix(task='multiclass', num_classes=n_seen_classes)

        for k, data_loader in data_loaders.items():
            # k is like 'camelyon16/val' or 'tcga_lung/test'
            cltor = self.eval_model(self.model, data_loader) # {'pred': {'y': all_gt, 'y_hat': all_pred, 'idx': all_idx, 'all_key_frequencies': all_key_frequencies}}

            for k_cltor, v_cltor in cltor.items():
                v_cltor['y'] += self.cfg['dataset_label_shift'][self.cfg['dataset_names'].index(k.split('/')[0])] # label shift
                if_binary = len(self.current_ensemble_classes['count']) == 2
                eval_results = self._eval_and_print(v_cltor, name=k+'/'+k_cltor, if_binary=if_binary)

                if self.current_dataset == self.cfg['dataset_names'][-1] and k.split('/')[-1] == 'test':
                    self.sum_test_acc += eval_results['acc@mid'] if if_binary else eval_results['acc']

                if k.split('/')[-1] == 'test':
                    acc = eval_results['acc@mid'] if if_binary else eval_results['acc']
                    acc = round(float(acc), 6)
                    current_test_acc.append(acc)
                    conf_mat_test.update(v_cltor['y_hat'].argmax(dim=1), v_cltor['y'])
                elif k.split('/')[-1] == 'val':
                    conf_mat_val.update(v_cltor['y_hat'].argmax(dim=1), v_cltor['y'])

                if self.current_dataset == self.cfg['dataset_names'][-1]:
                    if k.split('/')[-1] == 'val':
                        acc = eval_results['acc@mid'] if if_binary else eval_results['acc']
                        acc = round(float(acc), 6)
                        val_acc.append(acc)

                sids = self._get_unique_sids(k, v_cltor['idx'])
                path_save_pred = eval_result_dir + '/' + k.replace('/', '-') + '.csv'
                self._save_prediction_clf(sids, v_cltor['y'], v_cltor['y_hat'], path_save_pred, binary=if_binary)

        plot_confusion_matrix(conf_mat_val, save_dir=eval_result_dir, name='conf_matrix_val', idx=self.cfg['task_num'])
        plot_confusion_matrix(conf_mat_test, save_dir=eval_result_dir, name='conf_matrix_test', idx=self.cfg['task_num'])
        del conf_mat_test, conf_mat_val

        self.test_acc.append(current_test_acc)

        if self.current_dataset == self.cfg['dataset_names'][-1]:
            metrics_dir = osp.join(osp.dirname(self.current_save_result_dir), 'metrics')
            mkdir_if_missing(metrics_dir)

            with open(osp.join(metrics_dir, 'test_acc.txt'), 'w') as f:
                f.write(str(self.test_acc))

            with open(osp.join(metrics_dir, 'val_acc.txt'), 'w') as f:
                f.write(str(val_acc))

    def _get_unique_sids(self, k, idxs, concat=None):
        dataset_name = k.split('/')[0]
        dataset_split = k.split('/')[-1]
        sids = self.sids[dataset_name][dataset_split]
        idxs = idxs.tolist()
        if concat is None:
            return [sids[i] for i in idxs]
        else:
            return [sids[v] + "-" + str(concat[i].item()) for i, v in enumerate(idxs)]

    def _save_prediction_clf(self, sids, y_true, y_pred, save_path, binary=True, forgetting=False):
        if isinstance(y_true, Tensor):
            y_true = y_true.numpy()
        if isinstance(y_pred, Tensor):
            # pred is logit
            y_pred_raw = y_pred
            y_pred = softmax(y_pred, dim=1)  # apply softmax to logit outputs (N, num_cls)
            y_pred = y_pred.numpy()

        assert len(sids) == len(y_true)
        assert len(sids) == len(y_pred)

        save_data = {'sids': sids, 'y': y_true}
        cols = ['sids', 'y']
        if binary:
            save_data['y_hat'] = y_pred[:, 1]
            cols.append('y_hat')
        else:
            for i in range(y_pred.shape[-1]):
                _col = 'y_hat_' + str(i)
                save_data[_col] = y_pred[:, i]
                cols.append(_col)

        df = pd.DataFrame(save_data, columns=cols)
        df.to_csv(save_path, index=False)

        wandb_display_name = osp.splitext(osp.basename(save_path))[0].replace('-', '/')+"/dataframe"
        wandb.log({'forgetting_process/'+wandb_display_name if forgetting else wandb_display_name: wandb.Table(dataframe=df)})

        # add raw version
        save_data = {'sids': sids, 'y': y_true}
        cols = ['sids', 'y']
        for i in range(y_pred_raw.shape[-1]):
            _col = 'class_' + str(i)
            save_data[_col] = y_pred_raw[:, i]
            cols.append(_col) 
        df = pd.DataFrame(save_data, columns=cols)

        raw_save_path = save_path.replace('.csv', '_raw.csv')
        df.to_csv(raw_save_path, index=False)
        wandb_raw_display_name = osp.splitext(osp.basename(raw_save_path))[0].replace('-', '/')+"/dataframe"
        wandb.log({'forgetting_process/'+wandb_raw_display_name if forgetting else wandb_raw_display_name: wandb.Table(dataframe=df)})


    def _run_training(self):
        # iterative training
        last_epoch = -1

        for epoch in range(self.cfg['epochs'][self.cfg['task_num'] - 1]):
            last_epoch = epoch + 1
            print('[train] fold: {}, task: {}, epoch: {}'.format(self.cfg['data_split_seed'], self.cfg['task_num'], last_epoch))
            print('[train] lr={:.8f}'.format(self.optimizer.param_groups[0]['lr']))
            wandb.log({'train/lr/val': self.optimizer.param_groups[0]['lr']})

            # val datasets of previous tasks are also evaluated during training
            # supervise the catastrophic forgetting process
            if self.cfg['check_forgetting_process']:
                for k, v in self.data_loader.items():
                    if k == self.current_dataset:
                        break
                    eval_cltor = self.eval_model(self.model, v['val'])
                    for k_cltor, v_cltor in eval_cltor.items():
                        v_cltor['y'] += self.cfg['dataset_label_shift'][self.cfg['dataset_names'].index(k)]  # label shift
                        if_binary = len(self.current_ensemble_classes['count']) == 2
                        self._eval_and_print(v_cltor, name='forgetting_process/' + k + '/val/' + k_cltor, at_epoch=last_epoch, if_binary=if_binary)

                        sids = self._get_unique_sids(k + '/val', v_cltor['idx'])
                        forgetting_save_dir = osp.join(self.current_save_result_dir, 'forgetting_process')
                        mkdir_if_missing(forgetting_save_dir)
                        forgetting_save_path = osp.join(forgetting_save_dir, k + '-val.csv')
                        self._save_prediction_clf(sids, v_cltor['y'], v_cltor['y_hat'], forgetting_save_path, binary=if_binary, forgetting=True)

            train_cltor = self._train_each_epoch()

            if self.model_name == 'AlignMIL':
                # update tunable_v of old classes within their null spaces, Eq. 8
                tunable_v_delta = {}
                tun_lr = self.cfg['tunable_v_lr']
                print('[info] Tunable V update (lr):', tun_lr)
                with torch.no_grad():
                    for old_task in self.cfg["seen_task_dict"].keys()-{self.current_dataset}:
                        dataset_idx = self.cfg["seen_task_dict"][old_task]
                        for subtype in range(self.cfg['dataset_subtype_num'][dataset_idx]):
                            old_label = self.cfg['dataset_label_shift'][dataset_idx] + subtype
                            if self.phi[old_label].shape[1] > 0:
                                self.tunable_v[dataset_idx][subtype] = self.tunable_v[dataset_idx][subtype] + tun_lr * self.alpha[old_label] @ self.phi[old_label].t() # [1, C]
                                tunable_v_delta[old_label] = self.alpha[old_label] @ self.phi[old_label].t() # [1, C]
                print(f'[info] Tunable V update (norm):')
                for k, v in tunable_v_delta.items():
                    print(f'  {k}: {v.norm()}')
                print(f'[info] Tunable V update (abs max):')
                for k, v in tunable_v_delta.items():
                    print(f'  {k}: {v.abs().max()}')

            for k_cltor, v_cltor in train_cltor.items():
                self.train_key_frequency[self.current_dataset].append(v_cltor['all_key_frequencies']) # contains all the frequencies in every epochs, len = n_epochs
                if_binary = len(self.current_ensemble_classes['count']) == 2
                self._eval_and_print(v_cltor, name='train/'+k_cltor, at_epoch=last_epoch, if_binary=if_binary)
                # update the latest dataset2key_idx
                final_task_freq = self.train_key_frequency[self.current_dataset][-1]
                _, indices = torch.topk(final_task_freq, self.cfg['match_size'])
                self.dataset2key_idx[self.current_dataset] = indices

            eval_results = None
            eval_cltor = self.eval_model(self.model, self.data_loader[self.current_dataset]['val'])
            for k_cltor, v_cltor in eval_cltor.items():
                self.val_key_frequency[self.current_dataset].append(v_cltor['all_key_frequencies'])

                v_cltor['y'] += self.cfg['dataset_label_shift'][self.current_dataset_idx] # label shift
                if_binary = len(self.current_ensemble_classes['count']) == 2
                eval_results = self._eval_and_print(v_cltor, name='val/'+k_cltor, at_epoch=last_epoch, if_binary=if_binary)
                val_loss = eval_results['loss']
                val_acc = eval_results['acc@mid'] if if_binary else eval_results['acc']
                monitor_metric = (val_loss + (1 - val_acc)) / 2

            if last_epoch > self.cfg['lrs_warmup']:
                self.lr_scheduler.step(monitor_metric)
            self.early_stop(last_epoch, monitor_metric)

            if self.early_stop.save_ckpt():
                self._save_model(last_epoch, ckpt_type='best')
                print("[save best model] best model saved at epoch {}".format(last_epoch))

                if self.cfg['check_feature']:
                    save_dir = osp.join(self.current_save_result_dir, 'check_feature')
                    mkdir_if_missing(save_dir)
                    for k_cltor, v_cltor in train_cltor.items():
                        check_feature_distribution(v_cltor['all_feature_dict']['image'], v_cltor['y'], save_dir, type='image', epoch=last_epoch)
                        check_feature_distribution(v_cltor['all_feature_dict']['text'], v_cltor['y'], save_dir, type='text', epoch=last_epoch)
            if self.early_stop.stop():
                break

        self._save_model(last_epoch, ckpt_type='last')
        print("[save last model] last model saved at epoch {}".format(last_epoch))

        if self.cfg['check_key_matching']:
            print('[all key matching frequency on train dataset]', self.train_key_frequency[self.current_dataset])
            print('[all key matching frequency on validation dataset]', self.val_key_frequency[self.current_dataset])
            key_matching_save_dir = osp.join(self.current_save_result_dir, 'key_matching')
            mkdir_if_missing(key_matching_save_dir)
            plot_key_matching_heatmap(self.train_key_frequency[self.current_dataset], key_matching_save_dir, dataset='train')
            plot_key_matching_heatmap(self.val_key_frequency[self.current_dataset], key_matching_save_dir, dataset='val')
    
    def cal_null_space(self):
        print('[info] Calculate Null Space for current labels')
        loader = self.data_loader[self.current_dataset]['train']
        fb_dict = {}
        task_fb = []
        for data_idx, data_x, data_y in loader: # [B, 1], [B, N, C], [B, 1]: B = 1
            data_x = data_x.to(self.device)
            with torch.no_grad():
                data_x = self.base_model.visual.forward_project(data_x) # [B, N, C]

            # query vector
            fb, _ = torch.max(data_x, dim=1) # [B, C]

            label = data_y.item()
            if label not in fb_dict:
                fb_dict[label] = []
            fb_dict[label].append(fb)
            task_fb.append(fb)

        # class-level null space; phi is computed from the query matrix of the whole task
        current_labels = []
        for y in fb_dict.keys():
            label = y + self.cfg['dataset_label_shift'][self.current_dataset_idx]
            current_labels.append(label)
            fb = torch.cat(task_fb, dim=0)
            self.fb[label] = fb # [n_train, C]
            if self.current_dataset_idx > 0:
                phi = get_null_space(fb, a=10)
                if phi.shape[1] > 0:
                    print(f"phi {label} shape: {phi.shape}")
                    print(f"Norm of class {label} fb @ phi: {torch.norm(fb @ phi)}")
                self.phi[label] = phi # [C, rank]
                self.alpha[label] = nn.Parameter(0.02 * torch.randn(1, phi.shape[1], dtype=self.dtype, device=self.device)) # [1, rank]
            else:
                self.phi[label] = torch.empty([0, 0], device=self.device)
                self.alpha[label] = torch.empty([0, 0], device=self.device)

        if self.current_dataset_idx < 1:
            return
        # optimize alpha of current classes against query features of previous classes, Eq. 7
        print('[info] Update alpha by previous seen labels')
        lr = 0.01
        n_iteration = 100
        thr = self.cfg['class_feature_thr']
        print(f'[info] alpha update lr: {lr}, n_iteration: {n_iteration}, thr: {thr}')
        old_labels = []
        for old_task in self.cfg["seen_task_dict"].keys()-{self.current_dataset}:
            dataset_idx = self.cfg["seen_task_dict"][old_task]
            for subtype in range(self.cfg['dataset_subtype_num'][dataset_idx]):
                old_labels.append(self.cfg['dataset_label_shift'][dataset_idx] + subtype)

        param_list = []
        for _, alpha_param in self.alpha.items():
            if alpha_param.shape[1] > 0:
                param_list.append(alpha_param)

        if len(param_list) > 0:
            optimizer = torch.optim.Adam(param_list, lr=lr)

            for epo in range(n_iteration):
                optimizer.zero_grad()
                loss = 0.

                for cur_label in current_labels:
                    for old_label in old_labels:
                        res = self.fb[old_label] @ self.phi[cur_label] @ self.alpha[cur_label].t() + thr # [n_train, 1]
                        loss += torch.mean(relu(res))  # Calculate mean to get scalar loss

                # Only perform backward and optimization if we have parameters to update
                if len(param_list) > 0 and loss.item() > 1e-8:  # Add a small tolerance to avoid unnecessary backward on exact zero
                    loss.backward()
                    # Update parameters (alpha_t)
                    optimizer.step()
            
                # Print progress
                if (epo + 1) % 10 == 0:
                    print(f'Iteration [{epo + 1}/{n_iteration}], Loss: {loss.item():.4f}')

            print("alpha:", sorted(self.alpha.items()))

    def _save_model(self, epoch, ckpt_type='best'):
        net_ckpt_dict = self._get_state_dict(epoch)
        save_dir = osp.join(self.current_save_result_dir, 'model_ckpts')
        mkdir_if_missing(save_dir)
        torch.save(net_ckpt_dict, save_dir+f'/{ckpt_type}.pth')

    def _get_state_dict(self, epoch=None):
        if self.cfg['opt_name'] == 'adam':
            return_dict = {'epoch': epoch}
            return_dict.update({key: val for key, val in self.model.state_dict().items() if 'key' in key or 'prompt' in key or 'tunable_v' in key})
        else:
            raise NotImplementedError('Invalid optimizer')

        return return_dict

    @torch.no_grad()
    def eval_model(self, model, loader):
        model.eval()

        idx_collector, x_collector, y_collector = [], [], []
        all_idx, all_pred, all_gt, all_key_frequencies = [], [], [], torch.zeros(self.cfg['pool_size'], dtype=torch.int)
        i_batch = 0
        eval_every_batch = self.cfg['eval_every_batch']

        for data_idx, data_x, data_y in loader: # [B, 1], [B, N, C], [B, 1]: B = 1
            i_batch += 1

            data_x = data_x.to(self.device)

            # proj to VL space
            data_x = self.base_model.visual.forward_project(data_x) # [B, N, C]

            x_collector.append(data_x)
            y_collector.append(data_y)
            idx_collector.append(data_idx)

            # in a mini_batch
            if i_batch % eval_every_batch == 0:
                # evaluate network
                cur_pred, key_indices = model(x_collector, eval=True) # [MB, num_cls]
                all_pred.append(cur_pred.detach().cpu())
                all_gt.append(torch.cat(y_collector, dim=0))
                all_idx.append(torch.cat(idx_collector, dim=0))

                key_indices = key_indices.detach().cpu()
                key_frequencies = torch.zeros(self.cfg['pool_size'], dtype=torch.int)
                for index in key_indices.flatten():
                    key_frequencies[index] += 1
                all_key_frequencies += key_frequencies

                # reset mini-batch
                idx_collector, x_collector, y_collector = [], [], []
                torch.cuda.set_device(self.cfg['cuda_id'])
                torch.cuda.empty_cache()

        all_pred = torch.cat(all_pred, dim=0)  # [num_slide, num_cls]
        all_gt = torch.cat(all_gt, dim=0).squeeze(1)  # [num_slide, ]
        all_idx = torch.cat(all_idx, dim=0).squeeze(1)  # [num_slide, ]


        eval_cltor = dict()
        eval_cltor['pred'] = {'y': all_gt, 'y_hat': all_pred, 'idx': all_idx, 'all_key_frequencies': all_key_frequencies}

        return eval_cltor

    def _eval_and_print(self, cltor, name='', at_epoch=None, if_binary=True):
        if if_binary:
            eval_results = self.b_evaluator.compute(cltor, self.cfg['b_metrics_list'], self.cfg['loss_function'])
        else:
            eval_results = self.m_evaluator.compute(cltor, self.cfg['m_metrics_list'], self.cfg['loss_function'])
        wandb_eval_results = rename_keys(eval_results, name, sep='/')

        print("[{}] At epoch {}:".format(name, at_epoch), end=' ')
        print(' '.join(['{}={:.6f},'.format(k, v) for k, v in eval_results.items()]))
        wandb.log(wandb_eval_results)

        return eval_results

    def _train_each_epoch(self):
        self.model.eval()

        idx_collector, x_collector, y_collector = [], [], []
        all_pred, all_gt, all_key_frequencies, all_feature_dict = [], [], torch.zeros(self.cfg['pool_size'], dtype=torch.int), {'image': [], 'text': []}
        i_batch = 0
        train_loader = self.data_loader[self.current_dataset]['train']
        bp_every_batch = self.cfg['bp_every_batch']

        for data_idx, data_x, data_y in train_loader: # [B, 1], [B, N, C], [B, 1]: B = 1
            i_batch += 1

            data_y += self.cfg['dataset_label_shift'][self.current_dataset_idx] # label shift
            data_x = data_x.to(self.device)
            data_y = data_y.to(self.device)

            # proj to VL space
            with torch.no_grad():
                data_x = self.base_model.visual.forward_project(data_x) # [B, N, C]
            x_collector.append(data_x)
            y_collector.append(data_y)
            idx_collector.append(data_idx)

            # in a mini-batch
            if i_batch % bp_every_batch == 0:
                # update network
                cur_pred, key_indices, feature_dict = self._update_network(i_batch, x_collector, y_collector, idx_collector) # [MB, num_cls]
                all_pred.append(cur_pred.detach().cpu())
                all_gt.append(torch.cat(y_collector, dim=0).detach().cpu())

                key_indices = key_indices.detach().cpu()
                key_frequencies = torch.zeros(self.cfg['pool_size'], dtype=torch.int)
                for index in key_indices.flatten():
                    key_frequencies[index] += 1
                all_key_frequencies += key_frequencies
                if self.cfg['print_mb_key_matching']:
                    print('[key matching frequencies]', key_frequencies.tolist())

                all_feature_dict['image'].append(feature_dict['image'].detach().cpu())
                all_feature_dict['text'].append(feature_dict['text'].detach().cpu())

                # reset mini-batch
                idx_collector, x_collector, y_collector = [], [], []
                torch.cuda.set_device(self.cfg['cuda_id'])
                torch.cuda.empty_cache()

        all_pred = torch.cat(all_pred, dim=0)  # [num_slide, num_cls]
        all_gt = torch.cat(all_gt, dim=0).squeeze(1)  # [num_slide, ]
        all_feature_dict['image'] = torch.cat(all_feature_dict['image'], dim=0) # [num_slide, C]
        all_feature_dict['text'] = torch.cat(all_feature_dict['text'], dim=0) # [num_slide, num_cls, C]

        train_cltor = dict()
        train_cltor['pred'] = {'y': all_gt, 'y_hat': all_pred, 'all_key_frequencies': all_key_frequencies, 'all_feature_dict': all_feature_dict}

        return train_cltor

    def _update_network(self, i_batch, xs, ys, idx_num):
        """
        Update network using one mini-batch data
        """
        class_indices_dict = {}
        for dataset, idx in self.cfg['seen_task_dict'].items():
            class_st = self.cfg['dataset_label_shift'][idx]
            class_end = class_st + self.cfg['dataset_subtype_num'][idx]
            class_indices_dict[dataset] = list(range(class_st, class_end))

        if self.model_name == 'AlignMIL':
            bag_preds, loss_dict, key_indices, feature_dict = self.model(xs, current_dataset=self.current_dataset, class_indices=class_indices_dict) # [MB, num_cls]
        else:
            bag_preds, loss_dict, key_indices, feature_dict = self.model(xs) # [MB, num_cls]

        # zero gradients buffer
        self.optimizer.zero_grad()

        # classification loss
        bag_label = torch.cat(ys, dim=0).squeeze(-1)  # [MB, ]
        clf_loss = self.loss_function(bag_preds, bag_label)

        # total loss
        total_loss = clf_loss + self.cfg['lambda'] * loss_dict['matching_loss'] + self.cfg['beta'] * loss_dict['class_sim_loss']
        if self.model_name == 'AlignMIL' and len(self.dataset2key_idx) > 1 and len(self.dataset2key_idx) == len(class_indices_dict):
            pc_loss = self.model.compute_prototype_class_loss(self.dataset2key_idx, self.current_dataset, class_indices=class_indices_dict) # PCFA
            total_loss += pc_loss
            print(
                "[training epoch] {}-th batch: total loss = {:.6f}, classification loss = {:.6f}, matching loss = {:.6f}, class similarity loss = {:.6f}, prototype class loss = {:.6f}".format(
                    i_batch, total_loss.item(), clf_loss.item(), loss_dict['matching_loss'].item(),
                    loss_dict['class_sim_loss'].item(), pc_loss.item()))
        else:
            print(
                "[training epoch] {}-th batch: total loss = {:.6f}, classification loss = {:.6f}, matching loss = {:.6f}, class similarity loss = {:.6f}".format(
                    i_batch, total_loss.item(), clf_loss.item(), loss_dict['matching_loss'].item(), loss_dict['class_sim_loss'].item()))

        wandb_dict = {'train/total loss': total_loss.item(), 'train/classification loss': clf_loss.item(),
                      'train/matching loss': loss_dict['matching_loss'].item(),
                      'train/class similarity loss': loss_dict['class_sim_loss'].item()}
        wandb.log(wandb_dict)

        # backward gradients and update networks
        total_loss.backward()
        if self.cfg['max_norm'] != 'None':
            torch.nn.utils.clip_grad_norm_(self.model.parameters(), self.cfg['max_norm'])
        if self.cfg['opt_name'] == 'adam':
            for param in self.model.parameters():
                if param.grad is not None:
                    if torch.all(param.grad == 0):
                        param.grad = None # Only the matched key and prompt are updated.
        else:
            raise NotImplementedError('Invalid optimizer')
        self.optimizer.step()

        return bag_preds, key_indices, feature_dict